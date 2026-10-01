from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from telegram_business.products import (
    KNOWN_COLORS,
    _family_name,
    extract_product_query,
    normalize_model,
)


DEMAND_PATTERNS = (
    r"\bищ(?:у|ем|ет)\b", r"\bнуж(?:ен|на|но|ны)\b", r"\bкуплю\b",
    r"\bкто\s+(?:даст|продаст)\b", r"\bу\s+кого\s+есть\b",
    r"\bнадо\b", r"\bтребуется\b", r"\bесть\s+у\s+кого\b",
    r"\bkerak\b", r"\bkerek\b", r"\bkere\b", r"\bкерак\b",
    r"\bbor\s*mi\b", r"\bbormi\b", r"\bборми\b",
    r"\bkimda\s+bor\b", r"\bqidir", r"\bқидир",
    r"\bolaman\b", r"\bоламан\b", r"\btopib\s+ber", r"\bтопиб\s+бер",
)
SUPPLY_PATTERNS = (
    r"\bпродам\b", r"\bпродаю\b", r"\bв\s+наличии\b",
    r"\bесть\s+в\s+наличии\b", r"\bпредлагаю\b", r"\bимеется\b",
    r"\bsotaman\b", r"\bsotiladi\b", r"\bсотаман\b", r"\bсотилади\b",
    r"\bmavjud\b", r"\bмавжуд\b", r"\btayyor\b", r"\bтайёр\b",
)


@dataclass(frozen=True, slots=True)
class ModelMention:
    model_key: str
    model_name: str
    memory: str | None
    color: str | None
    confidence: float


@dataclass(frozen=True, slots=True)
class MessageAnalysis:
    intent: str
    mentions: tuple[ModelMention, ...]


def classify_intent(text: str) -> str:
    normalized = normalize_model(text)
    if any(re.search(pattern, normalized, re.IGNORECASE) for pattern in DEMAND_PATTERNS):
        return "demand"
    if any(re.search(pattern, normalized, re.IGNORECASE) for pattern in SUPPLY_PATTERNS):
        return "supply"
    return "unknown"


def _aliases(model_name: str) -> set[str]:
    normalized = normalize_model(model_name)
    aliases = {normalized}
    words = normalized.split()
    brands = {
        "apple", "iphone", "samsung", "galaxy", "xiaomi", "redmi", "poco",
        "honor", "huawei", "oppo", "vivo", "realme", "oneplus", "tecno",
        "infinix", "google", "pixel",
    }
    while words and words[0] in brands:
        words = words[1:]
        if len(words) >= 2 and any(any(c.isdigit() for c in word) for word in words):
            aliases.add(" ".join(words))
    return {alias for alias in aliases if alias}


def _canonical_model(value: str) -> str:
    """Collapse catalogue transport/radio variants into the searched family."""
    cleaned = re.sub(r"^\s*zzztex[.\s]+", "", str(value), flags=re.I)
    # Supplier exports may append ``(2)`` to duplicate rows.  It is not part of
    # the product name and otherwise makes a perfectly valid lookup ambiguous.
    cleaned = re.sub(r"\s+\(\d+\)\s*$", "", cleaned)
    cleaned = re.sub(r"\s+(?:4g|5g)(?:\s+\d+)?\s*$", "", cleaned, flags=re.I)
    family = str(_family_name(cleaned) or cleaned).strip()
    return family or str(value).strip()


def _message_forms(text: str) -> str:
    normalized = normalize_model(text)
    normalized = re.sub(r"\bgoogel\b", "google", normalized)
    normalized = re.sub(r"\b(whoop\s+5\s+0)\s+peek\b", r"\1 peak", normalized)
    normalized = re.sub(r"\b(air|pro)\s*(1[3456])\b", r"\1 \2", normalized)
    normalized = re.sub(r"\b(\d{1,2})\s*pm\b", r"\1 pro max", normalized)
    normalized = re.sub(r"\b(\d{1,2})\s*p\b", r"\1 pro", normalized)
    normalized = re.sub(r"\bs\s*(\d{1,2})\s*u\b", r"s\1 ultra", normalized)
    return " ".join(normalized.split())


def _special_aliases(model_name: str) -> set[str]:
    """Safe shop shorthand aliases used in the supplier group."""
    value = normalize_model(model_name)
    aliases: set[str] = set()
    iphone = re.fullmatch(
        r"(?:apple\s+)?iphone\s+(\d{1,2})(?:\s+(pro max|pro|plus|e))?",
        value,
    )
    if iphone:
        generation, variant = iphone.groups()
        if variant:
            aliases.add(f"{generation} {variant}")
        if variant == "pro max":
            aliases.update({f"{generation} max", f"{generation} pm"})
        return aliases
    samsung = re.fullmatch(
        r"samsung\s+galaxy\s+(?:tab\s+)?([asz]\d+(?:\s+(?:ultra|plus|fe))?)",
        value,
    )
    if samsung:
        aliases.add(samsung.group(1))
    macbook = re.fullmatch(
        r"(?:apple\s+)?(macbook\s+(?:air|pro)\s+\d+\s+m\d+)\s+\d+\s+core",
        value,
    )
    if macbook:
        aliases.add(macbook.group(1))
    return aliases


def _observed_mentions(text: str) -> tuple[ModelMention, ...]:
    """Capture explicit market models that are newer than the price catalogue."""
    value = _message_forms(text)
    _query, memory, color = extract_product_query(text)
    matches: list[tuple[str, str]] = []
    iphone = re.search(
        r"\b(?:iphone\s+)?(1[2-9]|2\d)\s*(pro\s+max|pro|max|plus|e)?\b",
        value,
    )
    phone_context = bool(
        re.search(r"\b(?:iphone|sim|esim)\b", value) or iphone and iphone.group(2)
    )
    if iphone and phone_context:
        generation, variant = iphone.groups()
        suffix = "Pro Max" if variant in {"pro max", "max"} else (
            variant.title() if variant else ""
        )
        name = f"Apple iPhone {generation}" + (f" {suffix}" if suffix else "")
        matches.append((normalize_model(name), name))
    airpods = re.search(r"\bairpods\s+(\d+)\b", value)
    if airpods:
        name = f"Apple AirPods {airpods.group(1)}"
        matches.append((normalize_model(name), name))
    ipad = re.search(r"\bipad\s+(air|pro|mini)?\s*(m\d+|\d+)\b", value)
    if ipad:
        family, generation = ipad.groups()
        name = "Apple iPad" + (f" {family.title()}" if family else "")
        name += f" {generation.upper()}"
        matches.append((normalize_model(name), name))
    if re.search(r"\b(?:redmi|mi)\s+pad\s+se\b", value):
        name = "Xiaomi Redmi Pad SE"
        matches.append((normalize_model(name), name))
    macbook = re.search(
        r"\bmacbook\s+(air|pro)\s+(1[3456])(?:\s+(1[3456]))?\s+(m\d+)\b",
        value,
    )
    if macbook:
        family, first_size, second_size, chip = macbook.groups()
        for size in (first_size, second_size):
            if size:
                name = f"Apple MacBook {family.title()} {size} {chip.upper()}"
                matches.append((normalize_model(name), name))
    if re.search(r"\boakley(?:\s+meta)?\s+hstn\b", value):
        name = "Oakley Meta HSTN"
        matches.append((normalize_model(name), name))
    watch = re.search(r"^\s*(\d{1,2})\s+(4[0-9])(?:\s*mm)?\b", value)
    if watch:
        series, size = watch.groups()
        name = f"Apple Watch Series {series}"
        matches.append((normalize_model(name), name))
    samsung = re.search(r"\b([az]\d{2})(?:\s+(ultra|plus|fe))?\b", value)
    if samsung:
        code, variant = samsung.groups()
        name = f"Samsung Galaxy {code.upper()}" + (
            f" {variant.upper() if variant == 'fe' else variant.title()}"
            if variant else ""
        )
        matches.append((normalize_model(name), name))
    unique = dict(matches)
    return tuple(
        ModelMention(
            key,
            name,
            f"{watch.group(2)}mm" if watch and name.startswith("Apple Watch") else memory,
            color,
            0.8,
        )
        for key, name in unique.items()
    )


class ProductModelIndex:
    def __init__(self, models: Iterable[str]):
        canonical: dict[str, str] = {}
        alias_map: dict[str, set[str]] = {}
        for raw in models:
            model = _canonical_model(str(raw or "").strip())
            key = normalize_model(model)
            if not key or key in {"model", "model name"}:
                continue
            canonical.setdefault(key, model)
            for alias in _aliases(model) | _special_aliases(model):
                alias_map.setdefault(alias, set()).add(key)
        self.models = canonical
        self.aliases = tuple(sorted(
            (
                (alias, next(iter(keys)))
                for alias, keys in alias_map.items()
                if len(keys) == 1
            ),
            key=lambda item: (-len(item[0].split()), -len(item[0]), item[0]),
        ))

    @classmethod
    def from_xlsx(cls, path: Path) -> "ProductModelIndex":
        if not path.is_file():
            raise FileNotFoundError(f"catalog not found: {path}")
        workbook = load_workbook(path, read_only=True, data_only=True)
        models: list[str] = []
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                columns = {
                    str(value or "").strip().casefold(): index
                    for index, value in enumerate(header)
                }
                model_index = next((
                    columns[name] for name in ("model", "model_name", "модель")
                    if name in columns
                ), None)
                if model_index is None:
                    continue
                for row in rows:
                    if model_index < len(row) and row[model_index]:
                        models.append(str(row[model_index]).strip())
        finally:
            workbook.close()
        if not models:
            raise ValueError("catalog contains no model column")
        return cls(models)

    def find(self, text: str, *, maximum: int = 12) -> tuple[ModelMention, ...]:
        haystack = f" {_message_forms(text)} "
        _query, memory, color = extract_product_query(text)
        found: dict[str, tuple[str, int]] = {}
        for alias, key in self.aliases:
            if f" {alias} " not in haystack:
                continue
            previous = found.get(key)
            score = len(alias.split())
            if previous is None or score > previous[1]:
                found[key] = (alias, score)
            if len(found) >= maximum * 3:
                break
        # Suppress a shorter catalogue title contained in a longer exact title,
        # for example iPhone 16 Pro when iPhone 16 Pro Max is present.
        selected: list[tuple[str, str, int]] = []
        ranked = sorted(
            ((key, alias, score) for key, (alias, score) in found.items()),
            key=lambda item: (-item[2], -len(item[1]), item[1]),
        )
        for key, alias, score in ranked:
            if any(f" {alias} " in f" {kept_alias} " for _, kept_alias, _ in selected):
                continue
            selected.append((key, alias, score))
            if len(selected) >= maximum:
                break
        color_value = color
        if color_value is None:
            normalized_words = _message_forms(text).split()
            color_value = next((word for word in normalized_words if word in KNOWN_COLORS), None)
        return tuple(
            ModelMention(
                model_key=key,
                model_name=self.models[key],
                memory=memory,
                color=color_value,
                confidence=1.0 if alias == key else 0.95,
            )
            for key, alias, _score in selected
        )


class MarketMessageAnalyzer:
    def __init__(self, index: ProductModelIndex):
        self.index = index

    def analyze(self, text: str) -> MessageAnalysis:
        intent = classify_intent(text)
        mentions = self.index.find(text)
        if intent == "demand" and not mentions:
            mentions = _observed_mentions(text)
        return MessageAnalysis(
            intent=intent,
            mentions=mentions,
        )
